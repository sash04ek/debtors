"""Формирование досудебной претензии в docx по образцу."""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from orgs import LETTER_BODY as LETTER_BODY_DEFAULT
from orgs import Organization

STREET_TYPES = {"ул": "ул.", "пер": "пер.", "пр": "пр.", "пр-т": "пр-т", "просп": "пр-т", "проезд": "проезд",
                "б-р": "б-р", "бул": "б-р", "пл": "пл.", "ш": "ш.", "наб": "наб.", "туп": "туп.",
                "площадка": "площадка", "мкр": "мкр."}


def split_address(addr: str) -> tuple[str, str]:
    """'Калинина ул 113' → ('ул. Калинина', '113'); '10-й пер 114' → ('пер. 10-й', '114')."""
    from orgs import clean_house_address
    parts = clean_house_address(addr).split()
    if len(parts) < 2:
        return str(addr or "").strip(), ""
    house, rest = parts[-1], parts[:-1]
    if len(parts) >= 4 and parts[-2].lower().strip(".") in ("корп", "корпус", "стр", "строение"):
        house, rest = " ".join(parts[-3:]), parts[:-3]                  # «19 корп 2» — одно целое
    for i, w in enumerate(rest):
        t = STREET_TYPES.get(w.lower().strip("."))
        if t:
            name = " ".join(rest[:i] + rest[i + 1:])
            return f"{t} {name}".strip(), house
    return " ".join(rest), house


def clean_flat(kv) -> str:
    if isinstance(kv, float) and kv.is_integer():
        kv = int(kv)
    return re.sub(r"['’`]+", "", str(kv or "")).strip()


def money(v: float) -> str:
    return f"{v:,.2f}".replace(",", " ").replace(".", ",")


def safe_name(s: str) -> str:
    return re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", s).strip(" ._")[:100] or "претензия"


def _runs(par, text: str, size: Pt):
    """Добавляет текст с **жирным**."""
    for i, chunk in enumerate(re.split(r"\*\*", text)):
        if chunk:
            r = par.add_run(chunk)
            r.bold = True if i % 2 else None
            r.font.size = size


def _no_borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    tblPr.append(b)


def build_claim(org: Organization, *, address: str, flat, debt: float, on_date: date,
                path: str | Path) -> None:
    street, house = split_address(address)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)

    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15

    # шапка: слева реквизиты организации, справа адресат
    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit = False
    _no_borders(tbl)
    for cell, w in zip(tbl.rows[0].cells, (Cm(9.2), Cm(0.4), Cm(7.4))):
        cell.width = w
    left, _, right = tbl.rows[0].cells
    left.paragraphs[0].text = ""
    for i, line in enumerate(org.header.splitlines() or [""]):
        p = left.paragraphs[0] if i == 0 else left.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        _runs(p, line, Pt(10))
    addressee = [f"Собственнику квартиры № {clean_flat(flat)}",
                 f"в многоквартирном доме № {house}", f"по {street}"]
    for i, line in enumerate(addressee):
        p = right.paragraphs[0] if i == 0 else right.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.3)
        _runs(p, line, Pt(11))

    doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _runs(t, "**ДОСУДЕБНАЯ ПРЕТЕНЗИЯ**", Pt(11))
    doc.add_paragraph()

    values = dict(org=org.name, agent=org.agent, agent_address=org.agent_address,
                  date=on_date.strftime("%d.%m.%Y"), debt=money(debt), days=org.days)
    body = org.body_template()
    for line in body.splitlines():
        if not line.strip():
            continue
        line = re.sub(r"\{(\w+)\}", lambda m: str(values.get(m.group(1), m.group(0))), line)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.25)
        _runs(p, line, Pt(11))

    doc.add_paragraph()
    doc.add_paragraph()
    p = doc.add_paragraph()
    _runs(p, "Дата: _______________" + " " * 40 + "подпись: __________________", Pt(11))
    doc.save(str(path))


def _flat_key(flat: str):
    m = re.match(r"\d+", flat)
    return (int(m.group()) if m else 10**9, flat)


def build_letter(org: Organization, *, items: list[tuple[str, object]], on_date: date,
                 path: str | Path) -> int:
    """Письмо в ЕИРЦ со списком адресов должников. items — (адрес дома из отчёта, квартира).
    Возвращает число адресов в письме (без повторов)."""
    from datetime import timedelta

    # дома — в порядке первого появления, квартиры внутри дома — по номеру
    by_house: dict[str, set[str]] = {}
    for addr, flat in items:
        by_house.setdefault(str(addr).strip(), set()).add(clean_flat(flat))
    lines = []
    for addr, flats in by_house.items():
        street, house = split_address(addr)
        for f in sorted(flats, key=_flat_key):
            lines.append(f"{org.city}, {street} {house}, кв. {f}")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(1), Cm(1.5)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.font.size = Pt(12)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.0

    def para(text="", *, align=None, size=Pt(12), bold=False, indent=None):
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        if indent is not None:
            p.paragraph_format.left_indent = indent
        _runs(p, f"**{text}**" if bold and text else text, size)
        return p

    doc.add_paragraph()
    for line in org.letter_header.splitlines():
        big = line.startswith("# ")
        para(line[2:] if big else line, align=WD_ALIGN_PARAGRAPH.CENTER,
             size=Pt(14) if big else Pt(12), bold=bool(line.strip()))
    doc.add_paragraph()
    for line in org.letter_to.splitlines():
        para(line, align=WD_ALIGN_PARAGRAPH.RIGHT)
    doc.add_paragraph()

    deadline = (on_date + timedelta(days=int(org.letter_days or 10))).strftime("%d.%m.%Y")
    body = (org.letter_body.strip() or LETTER_BODY_DEFAULT).replace("{date}", deadline)
    for chunk in body.splitlines():
        p = para(chunk.lstrip("\t"), align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.paragraph_format.space_after = Pt(6)
    for line in lines:
        para(line)

    doc.add_paragraph()
    doc.add_paragraph()
    para(org.sign_role, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=True)
    para(f"{org.name}   " + " " * 40 + "_____________ " + org.sign_name, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=True)
    doc.save(str(path))
    return len(lines)


def file_name(fio, address, flat) -> str:
    return safe_name(f"Претензия {fio} {address} кв {clean_flat(flat)}") + ".docx"
