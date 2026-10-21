import _isolate                      # DEBTORS_HOME во временную папку, до импорта storage
HOME = _isolate.HOME
import tempfile
import tkinter as tk
import unittest
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook

import app
import claimlog
import core
import extras
import owners
import storage

OWNER = {"fio": "Иванов Иван Иванович", "share": "1/2", "birth_date": "01.01.1980", "birth_place": "г. Город",
         "passport": "0000 000000", "reg_address": "ул. Тестовая, 1", "live_address": ""}
CARD = {"address": "Тестовая ул 1", "flat": "1", "owners": [OWNER, {"fio": "Иванова Анна Петровна", "birth_date": "02.02.1985"}],
        "cadastral": ""}


class ExtrasRowTest(unittest.TestCase):
    def test_no_data_gives_empty_cells(self):
        r = extras.row({"address": "Х ул 1", "flat": "1", "report_fio": "Иванов Иван"}, None, None, 40)
        self.assertEqual(len(r), len(extras.HEADERS))
        self.assertEqual(r, [""] * len(extras.HEADERS))

    def test_flags_and_personal_data(self):
        rec = {"sent": (date.today() - timedelta(days=50)).isoformat(), "closed": "", "court": ""}
        r = dict(zip(extras.HEADERS, extras.row({"address": "Тестовая ул 1", "flat": "1", "report_fio": "Иванов Иван"},
                                                CARD, rec, 40, "Судебный участок № 1")))
        self.assertEqual(r["Данные собственника"], "заполнены не полностью")       # у второго собственника нет паспорта и адреса
        self.assertEqual(r["Судебный участок"], "Судебный участок № 1")
        self.assertEqual(r["Ответ на претензию"], "срок ответа истёк")
        self.assertEqual(r["Претензия отправлена"], claimlog.long(rec["sent"]))
        self.assertEqual(r["Приказ подан в суд"], "")
        self.assertEqual(r["Собственник"], "1) Иванов Иван Иванович; 2) Иванова Анна Петровна")
        self.assertEqual(r["Паспорт"], "1) 0000 000000; 2) —")
        self.assertEqual(r["Адрес проживания"], "")                                    # ни у кого нет — пусто, а не «1) —; 2) —»

    def test_single_owner_plain_values_and_court_filing(self):
        card = dict(CARD, owners=[OWNER], cadastral="61:00:000:1")
        rec = {"sent": "2026-01-01", "closed": "2026-02-01", "court": "2026-03-01"}
        r = dict(zip(extras.HEADERS, extras.row({"address": "Тестовая ул 1", "flat": "1", "report_fio": ""}, card, rec, 40)))
        self.assertEqual((r["Паспорт"], r["Доля в праве"], r["Кадастровый номер"]), ("0000 000000", "1/2", "61:00:000:1"))
        self.assertEqual((r["Ответ на претензию"], r["Приказ подан в суд"]), ("ответ получен", "01.03.2026"))

    def test_headers_do_not_clash_with_report_columns(self):
        h = extras.unique_headers(["ФИО", "паспорт", "Долг"])
        self.assertIn("Паспорт (программа)", h)
        self.assertEqual(len(set(h)), len(h))


class ExportNameTest(unittest.TestCase):
    def test_org_in_default_name(self):
        self.assertEqual(app.export_file_name("ООО УО «Ромашка»", "Долг поквартирно 09.2026", 20),
                         "ООО УО Ромашка_Долг поквартирно 09.2026_топ20.xlsx")
        self.assertEqual(app.export_file_name("ООО УО «Ромашка»", "", 5), "ООО УО Ромашка_должники.xlsx")

    def test_org_not_duplicated_and_name_is_safe(self):
        self.assertEqual(app.export_file_name("ООО УО «Ромашка»", "ООО УО Ромашка_отчёт_топ20", 20),
                         "ООО УО Ромашка_отчёт_топ20_топ20.xlsx")
        name = app.export_file_name('ООО "А/Б: очень длинное название управляющей компании ' + "я" * 80 + '"', "файл", 3)
        self.assertNotRegex(name, r'[\\/:*?"<>|]')
        self.assertLessEqual(len(name), 120 + len(".xlsx"))
        self.assertEqual(app.export_file_name("", "файл", 3), "файл_топ3.xlsx")


class ExportFileTest(unittest.TestCase):
    def test_extra_columns_written_and_dropped_on_reading(self):
        s = core.Settings(top_n=20, name_col="ФИО", debt_col="Долг", addr_col="Адрес", flat_col="Кв")
        res = core.process(core.Sheet(["Адрес", "Кв", "ФИО", "Долг"], [["Тестовая ул 1", "1", "Иванов Иван", 500.0],
                                                                       ["Тестовая ул 2", "2", "Петров Пётр", 300.0]]), s)
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "top.xlsx"
            core.export(res, p, "Долг", {}, ["Флаг"], [["да"], []])
            ws = load_workbook(p)["Топ-2 должников"]
            self.assertEqual([c.value for c in ws[1]], ["№", "Адрес", "Кв", "ФИО", "Долг", "Флаг"])
            self.assertEqual([c.value for c in ws[2]][-1], "да")
            self.assertIsNone([c.value for c in ws[3]][-1])
            back = core.read_filtered(p)
            self.assertEqual(back["headers"], ["Адрес", "Кв", "ФИО", "Долг"])
            self.assertEqual(len(back["rows"]), 2)
            self.assertEqual(len(back["rows"][0]), 4)
            core.export(res, p, "Долг", {})                                              # без доп. колонок — как раньше
            self.assertEqual(core.read_filtered(p)["headers"], ["Адрес", "Кв", "ФИО", "Долг"])


class AppExportTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        saved = (storage.CLAIMS_PATH, app.save_settings, app.messagebox.showinfo, app.messagebox.showerror,
                 app.filedialog.asksaveasfilename, owners._load_all)

        def restore():
            (storage.CLAIMS_PATH, app.save_settings, app.messagebox.showinfo, app.messagebox.showerror,
             app.filedialog.asksaveasfilename, owners._load_all) = saved
        self.addCleanup(self.tmp.cleanup)
        self.addCleanup(restore)
        storage.CLAIMS_PATH = Path(self.tmp.name) / "claims.json"
        app.save_settings = lambda s: None
        self.errors = []
        self.infos = []
        app.messagebox.showinfo = lambda *a, **kw: self.infos.append(a[1])
        app.messagebox.showerror = lambda *a, **kw: self.errors.append(a)
        self.out = Path(self.tmp.name) / "out.xlsx"
        self.asked = []
        app.filedialog.asksaveasfilename = lambda **kw: (self.asked.append(kw), str(self.out))[1]
        owners._load_all = lambda: {owners.make_key("Тестовая ул 1", "1"): CARD}
        try:
            self.app = app.App()
        except tk.TclError:
            self.skipTest("нет дисплея")
        self.addCleanup(self.app.destroy)

    def test_export_has_flags_and_personal_data_for_each_address(self):
        a, s = self.app, self.app.settings
        s.addr_col, s.flat_col, s.name_col, s.debt_col = "Адрес", "Кв", "ФИО", "Долг"
        for key, val in (("addr_col", "Адрес"), ("flat_col", "Кв"), ("name_col", "ФИО"), ("debt_col", "Долг")):
            a.col_vars[key].set(val)
        h = ["ФИО", "Адрес", "Кв", "Долг"]
        rows = [["Иванов Иван", "Тестовая ул 1", "1", 500.0], ["Петров Пётр", "Тестовая ул 1", "2", 300.0]]
        a.result = core.Result(h, rows, [500.0, 300.0], {})
        a.show_rows(h, rows, numbered=True)
        claimlog.mark_sent([("Тестовая ул 1", "1")], date.today() - timedelta(days=3))
        a.export()
        self.assertEqual(self.errors, [])
        self.assertIn(a.current_org().name.replace("«", "").replace("»", ""), self.asked[-1]["initialfile"])    # УК в имени файла
        self.assertIn("персональные данные", self.infos[-1])                           # предупреждение: в файле паспорта и адреса
        ws = load_workbook(self.out)["Топ-2 должников"]
        head = [c.value for c in ws[1]]
        for name in extras.HEADERS:
            self.assertIn(name, head)
        r1, r2 = ([c.value for c in ws[i]] for i in (2, 3))
        col = lambda r, name: r[head.index(name)]
        self.assertEqual(col(r1, "Данные собственника"), "заполнены не полностью")
        self.assertEqual(col(r1, "Ответ на претензию"), "ждём ответа")
        self.assertEqual(col(r1, "Паспорт"), "1) 0000 000000; 2) —")
        for name in extras.HEADERS:                                                      # адрес без карточки и журнала — всё пусто
            self.assertIn(col(r2, name), (None, ""), name)
        self.assertEqual(core.read_filtered(self.out)["headers"], h)                     # при повторном открытии колонки не множатся


if __name__ == "__main__":
    unittest.main()
