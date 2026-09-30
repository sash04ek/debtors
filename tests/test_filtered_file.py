"""Экспорт и повторное открытие отфильтрованного файла: python -m unittest discover tests"""
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import core  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

HEAD = ["Адрес", "Кв", "ФИО", "Долг", "Дата посл опл"]
ROWS = [
    ["Бабушкина ул 47", "10", "Иванов Иван Иванович", 500.5, "01.03.2020"],
    ["Абрикосовая ул 3", "1", "Сидорова Анна Сергеевна", 700.0, ""],
    ["Чехова ул 322", "5", "Кузнецов Олег Павлович", 100.0, "10.12.2021"],
    ["Итого по дому", None, None, 1300.5, None],             # строка без ФИО (подытог) — в отбор не попадает
]


def make_result():
    s = core.Settings(top_n=20, name_col="ФИО", debt_col="Долг", addr_col="Адрес", flat_col="Кв")
    return core.process(core.Sheet(HEAD, [list(r) for r in ROWS]), s)


class FilteredFileTest(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())
        self.path = self.dir / "top.xlsx"

    def test_export_marks_file_and_it_roundtrips(self):
        res = make_result()
        core.export(res, self.path, "Долг", meta={"org": "ООО Тест", "columns": {"name_col": "ФИО"}})
        kind, meta = core.detect_file_kind(self.path)
        self.assertEqual(kind, "filtered")
        self.assertEqual(meta["org"], "ООО Тест")
        self.assertEqual(meta["debt_col"], "Долг")
        # служебный лист скрыт, видимые листы — список и статистика
        wb = load_workbook(self.path)
        self.assertEqual(wb[core.META_SHEET].sheet_state, "hidden")
        self.assertEqual([n for n in wb.sheetnames if wb[n].sheet_state == "visible"], [f"Топ-3 должников", "Статистика"])
        data = core.read_filtered(self.path)
        self.assertEqual(data["headers"], HEAD)                                  # колонка «№» не попала в данные
        self.assertEqual(len(data["rows"]), 3)                                   # строка «Итого» отброшена
        self.assertEqual([r[2] for r in data["rows"]], [r[2] for r in res.top])
        self.assertEqual(data["amounts"], res.amounts)
        self.assertEqual(data["stats"]["отобрано"], 3)
        self.assertEqual(data["meta"]["org"], "ООО Тест")

    def test_original_files_are_recognised_as_original(self):
        wb = Workbook()
        ws = wb.active
        ws.append(HEAD)
        for r in ROWS:
            ws.append(r)
        wb.save(self.path)
        self.assertEqual(core.detect_file_kind(self.path), ("original", {}))
        self.assertEqual(core.detect_file_kind(self.dir / "old.xls")[0], "original")     # .xls всегда оригинал
        with self.assertRaises(ValueError):
            core.read_filtered(self.path)

    def test_file_from_older_version_without_meta_sheet(self):
        core.export(make_result(), self.path, "Долг")
        wb = load_workbook(self.path)
        del wb[core.META_SHEET]                                                   # как файл прежней версии программы
        wb.save(self.path)
        kind, meta = core.detect_file_kind(self.path)
        self.assertEqual((kind, meta), ("filtered", {"legacy": True}))
        data = core.read_filtered(self.path)
        self.assertEqual(data["meta"]["debt_col"], "Долг")                        # колонка долга определена по названию
        self.assertEqual(len(data["rows"]), 3)
        self.assertAlmostEqual(sum(data["amounts"]), 1300.5)

    def test_damaged_meta_falls_back_to_structure(self):
        core.export(make_result(), self.path, "Долг")
        wb = load_workbook(self.path)
        wb[core.META_SHEET]["A1"] = "не json"
        wb.save(self.path)
        self.assertEqual(core.detect_file_kind(self.path)[0], "filtered")


if __name__ == "__main__":
    unittest.main()
